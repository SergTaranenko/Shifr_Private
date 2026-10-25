"""Прогноз из модели: справедливая цена, зоны, из чего она складывается, допущения, тезисы.

Источник — Excel-книга модели, которую мы ведём вместе с Claude и храним в GitHub (папка data/).
Бот модель не меняет: он только читает из книги готовые цифры.
"""
from __future__ import annotations

import hashlib
import io
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from .config import MSK

PARAM_LABELS = {
    "usdrub": "Курс USD/RUB (прогноз на горизонт оценки)",
    "gold_usd": "Цена золота (прогноз на горизонт оценки)",
    "key_rate_2027": "Средняя ключевая ставка 2027 (прогноз ЦБ)",
    "ofz10": "Доходность 10-летних ОФЗ (для сравнения)",
    "p_peace": "Мир: вероятность мира и снятия всех санкций",
    "t_peace": "Мир: лет до снятия санкций",
    "buy_disc": "Зона покупки: скидка к справедливой",
    "sell_prem": "Цена продажи: премия к справедливой",
    "ai_rf": "ИИ: поправка на стоимость денег в РФ",
    "ai_mkt": "ИИ: рынок генеративного ИИ в РФ в 2030",
}
MODEL_NAMES = ("model.xlsx", "Модели_справедливой_стоимости.xlsx")
YELLOW = ("FFFFFF00", "00FFFF00")


class ModelFileError(Exception):
    """Книга не подходит: нет нужных листов или не посчитаны формулы."""


def _num(v):
    return float(v) if isinstance(v, (int, float)) else None


def fmt_value(v, number_format: str, unit: str) -> str:
    """Значение вводной так, как его видно в Excel: 8%, 4,5x, 1 950 млрд ₽."""
    if not isinstance(v, (int, float)):
        return str(v)
    nf = number_format or ""
    if unit == "%" or "%" in nf:
        s = f"{v * 100:.1f}".rstrip("0").rstrip(".") + "%"
    elif unit == "x" or '"x"' in nf or "\\x" in nf:
        s = f"{v:.1f}x"
    else:
        s = f"{v:,.0f}" if float(v).is_integer() or abs(v) >= 100 else f"{v:,.2f}".rstrip("0").rstrip(".")
        s = s.replace(",", " ")
    s = s.replace(".", ",")
    if unit and unit not in ("%", "x"):
        s += f" {unit}"
    return s


def short_component(label: str) -> str:
    if label.startswith("ИИ"):
        return "ИИ"
    if "после мира" in label:
        return "«Северные потоки» после мира"
    if label.startswith("Мир"):
        return "мир и снятие санкций"
    if "Северные потоки" in label:
        return "«Северные потоки»"
    if "Сила Байкала" in label:
        return "«Сила Байкала»"
    return label.split(":")[0]


def parse_workbook(data: bytes | str | Path) -> dict:
    """Читает «Сводку», «Параметры» и листы эмитентов. Возвращает прогноз в виде словаря."""
    raw = bytes(data) if isinstance(data, (bytes, bytearray)) else Path(data).read_bytes()
    try:
        wb = load_workbook(io.BytesIO(raw), data_only=True)
    except Exception as e:  # noqa: BLE001
        raise ModelFileError(f"не удалось открыть файл как Excel ({e.__class__.__name__})") from e
    if "Сводка" not in wb.sheetnames:
        raise ModelFileError("в книге нет листа «Сводка»")

    ws = wb["Сводка"]
    issuers = []
    for r in range(5, ws.max_row + 1):
        t = ws.cell(row=r, column=1).value
        if not t or not isinstance(t, str) or t.startswith("Итого") or len(t) > 6:
            if issuers:
                break
            continue
        row = {
            "ticker": t.strip(),
            "name": ws.cell(row=r, column=2).value,
            "method": ws.cell(row=r, column=3).value,
            "model_price": _num(ws.cell(row=r, column=4).value),
            "fair": _num(ws.cell(row=r, column=5).value),
            "buy": _num(ws.cell(row=r, column=6).value),
            "near": _num(ws.cell(row=r, column=7).value),
            "sell": _num(ws.cell(row=r, column=8).value),
            "qty": int(_num(ws.cell(row=r, column=12).value) or 0),
            "fair_no_peace": _num(ws.cell(row=r, column=15).value),
        }
        if None in (row["fair"], row["buy"], row["near"], row["sell"]):
            raise ModelFileError(
                f"для {t} в «Сводке» нет посчитанных значений. Откройте книгу в Excel, дождитесь пересчёта и сохраните её ещё раз"
            )
        issuers.append(row)
    if not issuers:
        raise ModelFileError("в «Сводке» не найдено ни одной бумаги")

    for it in issuers:
        it.update(_parse_issuer_sheet(wb, it["ticker"]))
        comp_sum = sum(v for _, v in it["components"])
        it["base"] = it["fair"] - comp_sum if it["components"] else it["fair"]

    params = {}
    if "Параметры" in wb.sheetnames:
        wp = wb["Параметры"]
        by_label = {wp.cell(row=r, column=1).value: wp.cell(row=r, column=2).value for r in range(1, wp.max_row + 1)}
        for key, label in PARAM_LABELS.items():
            v = _num(by_label.get(label))
            if v is not None:
                params[key] = v

    return {
        "version": hashlib.sha1(raw).hexdigest()[:12],
        "updated": datetime.now(MSK).strftime("%d.%m.%Y %H:%M"),
        "assumptions": params,
        "issuers": issuers,
    }


def _parse_issuer_sheet(wb, ticker: str) -> dict:
    out = {"thesis": "", "breaks": [], "watch": [], "levers": [], "components": []}
    if ticker not in wb.sheetnames:
        return out
    ws = wb[ticker]
    section = None
    for r in range(1, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        if not isinstance(a, str):
            continue
        s = a.strip()
        if s == "ВВОДНЫЕ":
            section = "inputs"
            continue
        if s == "РАСЧЁТ":
            section = "calcs"
            continue
        if s.startswith("ТЕЗИС И УСЛОВИЯ ЕГО СЛОМА"):
            section = "breaks"
            continue
        if s.startswith("ЧТО ОТСЛЕЖИВАЕТ БОТ"):
            section = "watch"
            continue
        if s in ("РЕЗУЛЬТАТ", "ПЕРЕКРЁСТНАЯ ПРОВЕРКА", "ИСТОЧНИКИ") or s.startswith("ЧУВСТВИТЕЛЬНОСТЬ"):
            section = None
            continue
        b = ws.cell(row=r, column=2)
        if section == "inputs":
            rgb = getattr(b.fill.fgColor, "rgb", None) if b.fill is not None else None
            if b.fill is not None and b.fill.fill_type == "solid" and rgb in YELLOW and b.value is not None:
                unit = ws.cell(row=r, column=3).value or ""
                out["levers"].append([s, fmt_value(b.value, b.number_format, unit)])
        elif section == "calcs":
            unit = ws.cell(row=r, column=3).value or ""
            if unit == "₽" and isinstance(b.value, (int, float)) and ("прибавка" in s or "вклад" in s):
                out["components"].append([short_component(s), float(b.value)])
        elif section == "breaks" and s.startswith("Тезис:"):
            out["thesis"] = s.removeprefix("Тезис:").strip()
        elif section in ("breaks", "watch") and s.startswith("•"):
            out[section].append(s.lstrip("• ").strip())
    return out


# --------------------------------------------------------------- где лежит модель
def find_model_file(folder: Path) -> Path | None:
    """Книга модели в папке data/: model.xlsx, «Модели_справедливой_стоимости.xlsx» или единственный .xlsx."""
    for name in MODEL_NAMES:
        p = folder / name
        if p.exists():
            return p
    files = sorted(folder.glob("*.xlsx"))
    return files[0] if len(files) == 1 else None


def load_bundled(bundled: Path) -> dict:
    """Модель из репозитория: сначала Excel-книга, если её нет или она не читается — levels.json."""
    xlsx = find_model_file(bundled)
    if xlsx is not None:
        try:
            lv = parse_workbook(xlsx)
            lv["source"] = xlsx.name
            lv["_xlsx"] = str(xlsx)
            lv["_origin"] = "github"
            return lv
        except ModelFileError as e:
            fallback_error = str(e)
        else:
            fallback_error = ""
    else:
        fallback_error = ""
    p = bundled / "levels.json"
    if not p.exists():
        raise FileNotFoundError("в data/ нет ни книги модели, ни levels.json")
    lv = json.loads(p.read_text(encoding="utf-8"))
    lv.setdefault("version", hashlib.sha1(p.read_bytes()).hexdigest()[:12])
    lv["_origin"] = "github"
    if fallback_error:
        lv["_warning"] = f"книгу в репозитории прочитать не удалось ({fallback_error}), взял levels.json"
    return lv


def load_override(data_dir: Path) -> dict | None:
    """Модель, присланная в чат (живёт на диске хостинга, пока в GitHub не появится новая)."""
    p = data_dir / "levels.json"
    if not p.exists():
        return None
    lv = json.loads(p.read_text(encoding="utf-8"))
    lv["_origin"] = "chat"
    x = data_dir / "model.xlsx"
    if x.exists():
        lv["_xlsx"] = str(x)
    return lv


def save_override(levels: dict, data: bytes, data_dir: Path) -> None:
    clean = {k: v for k, v in levels.items() if not k.startswith("_")}
    (data_dir / "levels.json").write_text(json.dumps(clean, ensure_ascii=False, indent=1), encoding="utf-8")
    (data_dir / "model.xlsx").write_bytes(data)


def drop_override(data_dir: Path) -> None:
    for name in ("levels.json", "model.xlsx"):
        p = data_dir / name
        if p.exists():
            p.unlink()


def snapshot(levels: dict) -> dict:
    """Компактная копия прогноза для сравнения со следующей версией."""
    return {
        "version": levels.get("version"),
        "updated": levels.get("updated"),
        "issuers": [{"ticker": i["ticker"], "fair": i["fair"]} for i in levels.get("issuers", [])],
        "assumptions": levels.get("assumptions", {}),
    }


def load_aliases(bundled: Path) -> dict:
    p = bundled / "aliases.json"
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else {}


def diff_levels(old: dict, new: dict) -> list[str]:
    """Короткий список изменений справедливых цен."""
    old_by = {i["ticker"]: i for i in old.get("issuers", [])}
    lines = []
    for it in new["issuers"]:
        o = old_by.get(it["ticker"])
        if not o:
            lines.append(f"{it['ticker']}: новая бумага, справедливая {it['fair']:,.0f} ₽".replace(",", " "))
            continue
        if abs((o.get("fair") or 0) - it["fair"]) >= 0.01:
            ch = it["fair"] / o["fair"] - 1 if o.get("fair") else 0
            lines.append(f"{it['ticker']}: {o['fair']:,.0f} → {it['fair']:,.0f} ₽ ({ch:+.1%})".replace(",", " "))
    for t in old_by:
        if t not in {i["ticker"] for i in new["issuers"]}:
            lines.append(f"{t}: убрана из модели")
    return lines
